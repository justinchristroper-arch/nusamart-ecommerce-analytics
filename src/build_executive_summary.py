"""
build_executive_summary.py
Membuat executive summary 7 slide untuk manajemen NusaMart: docs/executive_summary.pptx.

Semua angka dibaca langsung dari results/ (hasil query sql/02-05 dan
src/right_censoring_check.py), sehingga deck selalu sama dengan SQL dan dashboard.
Badan slide memuat paling banyak 3 angka kunci; angka pendukung dan file sumbernya
ada di speaker notes.

Versi PDF (docs/executive_summary.pdf) diekspor dari PowerPoint:
buka file PPTX, lalu File > Export > Create PDF/XPS.

Cara pakai (dari folder root repo):
    .venv\\Scripts\\python.exe src/build_executive_summary.py
"""
from datetime import datetime
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
RESULTS = ROOT / "results"
OUT = ROOT / "docs" / "executive_summary.pptx"

DASHBOARD_URL = (
    "https://app.powerbi.com/view?r=eyJrIjoiNWQ3ZjljMjgtZmRjZS00YjhjLWI2ZjctZGViNmIxNTJiMTJiIiwidCI6IjM0ODViOTYzLTgy"
    "YmEtNGE2Zi04MTBmLWI1Y2MyMjZmZjg5OCIsImMiOjEwfQ%3D%3D"
)
REPO_URL = "https://github.com/justinchristroper-arch/nusamart-ecommerce-analytics"
AUTHOR = "Justin Christroper"
PREPARED_ON = "September 2026"
MIN_STATE_ORDERS = 1000  # state dengan order lebih sedikit angkanya mudah berfluktuasi
NEAR_AVG_PP = 0.25  # selisih freight burden sekecil ini dari rata-rata nasional disebut "hampir sama"
PILOT_STATES = ["BA", "PE"]  # kandidat pilot subsidi ongkir (docs/insights_and_recommendations.md, Insight 5)
N_SLIDES = 7

# Warna mengikuti dashboard Power BI
INK = RGBColor(0x1F, 0x23, 0x28)
MUTED = RGBColor(0x5F, 0x6B, 0x7A)
BLUE = RGBColor(0x11, 0x8D, 0xFF)  # revenue / fokus utama
ORANGE = RGBColor(0xE6, 0x6C, 0x37)  # ongkir
LIGHT_BLUE = RGBColor(0xB3, 0xD7, 0xFF)
LIGHT_ORANGE = RGBColor(0xF5, 0xC6, 0xB0)
GREY = RGBColor(0x8C, 0x95, 0x9F)
GRID = RGBColor(0xD0, 0xD7, 0xDE)
PANEL = RGBColor(0xF3, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Segoe UI"
FONT_BOLD = "Segoe UI Semibold"


# ---------- format angka (gaya Indonesia, sama dengan README) ----------
def idn(value, decimals=0):
    """12342450.49 -> '12.342.450,49'"""
    text = f"{value:,.{decimals}f}"
    return text.replace(",", "_").replace(".", ",").replace("_", ".")


def pct(value, decimals=2):
    return f"{idn(value, decimals)}%"


def rp(value, decimals=2):
    return f"R$ {idn(value, decimals)}"


def growth(new, old):
    return (new / old - 1) * 100


# ---------- angka dari results/ ----------
def read(name):
    return pd.read_csv(RESULTS / f"{name}.csv")


def load_facts():
    kpi = read("05_kpi_total_filtered").iloc[0]
    q1c = read("02_q1c_fair_period_comparison").set_index("year")
    q1b = read("02_q1b_yoy").set_index("month")
    q4a = read("03_q4a_repeat_rate").iloc[0]
    q4b = read("03_q4b_order_distribution").set_index("n_orders")
    q4c = read("03_q4c_new_vs_returning")
    q4d = read("03_q4d_cohort_12m_repeat")
    q5a = read("03_q5a_decile").set_index("decile")
    q5c = read("03_q5c_decile_repeat").set_index("decile")
    q3 = read("02_q3_region").set_index("customer_state")
    q7 = read("04_q7_high_burden_categories").sort_values("freight_burden_pct", ascending=False)
    q6 = read("04_q6_volume_vs_value").set_index("category_en")
    q2c = read("02_q2c_category_fair_period").set_index("category_en")

    def returning_share(df):
        return df["returning_revenue"].sum() / (df["new_revenue"].sum() + df["returning_revenue"].sum()) * 100

    avg_price = kpi["total_revenue"] / kpi["units_sold"]
    office = q6.loc["office_furniture"]
    cheap = q7[q7["category_en"] != "office_furniture"]
    big = q3[q3["orders"] >= MIN_STATE_ORDERS]
    f = {
        "kpi": kpi,
        "q1c": q1c,
        "growth": {c: growth(q1c.loc[2018, c], q1c.loc[2017, c]) for c in ["revenue", "orders", "customers", "aov"]},
        "yoy_jan": q1b.loc["2018-01-01", "yoy_pct"],
        "yoy_jul": q1b.loc["2018-07-01", "yoy_pct"],
        "repeat": q4a,
        "one_time": q4b.loc[1],
        "cohort": q4d[q4d["cohort"].str.startswith("TOTAL")].iloc[0],
        "returning_share": returning_share(q4c),
        "returning_share_2018": returning_share(q4c[q4c["month"] >= "2018-01-01"]),
        "decile_share": q5a["revenue_share_pct"],
        "top20": q5a.loc[2, "cumulative_share_pct"],
        "bottom_half": 100 - q5a.loc[5, "cumulative_share_pct"],
        "d1": q5c.loc[1],
        "q3": q3,
        "big_states": big.sort_values("freight_burden_pct"),
        "small_states": q3[q3["orders"] < MIN_STATE_ORDERS].sort_values("freight_burden_pct", ascending=False),
        "big_order_share": big["orders"].sum() / q3["orders"].sum() * 100,
        "q7": q7,
        "q7_revenue": q7["revenue"].sum(),
        "q7_share": q7["revenue"].sum() / kpi["total_revenue"] * 100,
        "cheap": cheap,
        "cheap_below_avg": int((q6.loc[cheap["category_en"], "avg_item_price"] < avg_price).sum()),
        "avg_price": avg_price,
        "freight_unit_avg": kpi["freight_burden_pct"] / 100 * avg_price,
        "freight_unit_office": office["freight_burden_pct"] / 100 * office["avg_item_price"],
        "units_per_order": kpi["units_sold"] / kpi["total_orders"],
        "outside_sp": 100 - q3.loc["SP", "revenue_share_pct"],
        "sp_order_share": q3.loc["SP", "orders"] / kpi["total_orders"] * 100,
        "q2c": q2c,
    }
    # Narasi slide bergantung pada fakta ini; kalau data berubah dan cek gagal, teks slide harus ditinjau ulang
    assert q3["freight_burden_pct"].idxmin() == "SP", "SP bukan lagi state dengan freight burden terendah"
    assert abs(f["growth"]["aov"]) < 1, "AOV tidak lagi datar"
    assert f["cheap_below_avg"] == len(cheap), "ada kategori Q7 (selain office_furniture) yang tidak murah"
    assert f["freight_unit_office"] > 1.5 * f["freight_unit_avg"], "ongkir office_furniture tidak lagi jauh di atas rata-rata"
    return f


# ---------- bantuan menggambar ----------
def R(text, size=14, color=INK, font=FONT, italic=False, link=None):
    """Satu potongan teks (run) beserta gayanya."""
    return {"text": text, "size": size, "color": color, "font": font, "italic": italic, "link": link}


def P(*runs, after=0, align=PP_ALIGN.LEFT, bullet=False, line=None):
    """Satu paragraf berisi satu atau beberapa run."""
    return {"runs": runs, "after": after, "align": align, "bullet": bullet, "line": line}


def add_run(paragraph, spec):
    run = paragraph.add_run()
    run.text = spec["text"]
    font = run.font
    font.name = spec["font"]
    font.size = Pt(spec["size"])
    font.italic = spec["italic"]
    font.color.rgb = spec["color"]
    font.language_id = MSO_LANGUAGE_ID.INDONESIAN
    if spec["link"]:
        run.hyperlink.address = spec["link"]


def add_bullet(paragraph):
    """Bullet dengan indentasi gantung (python-pptx belum punya API untuk ini)."""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.22)))
    pPr.set("indent", str(-Inches(0.22)))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))


def add_text(slide, name, box, paragraphs, anchor=MSO_ANCHOR.TOP):
    x, y, w, h = box
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, spec in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = spec["align"]
        if spec["line"]:
            p.line_spacing = spec["line"]
        p.space_after = Pt(spec["after"])
        if spec["bullet"]:
            add_bullet(p)
        for run in spec["runs"]:
            add_run(p, run)
    return shape


def add_box(slide, name, box, fill=None, line=None, dashed=False, shape=MSO_SHAPE.RECTANGLE):
    x, y, w, h = box
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.name = name
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
        if dashed:
            s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return s


def add_line(slide, name, x1, y, x2, color=GRID):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y), Inches(x2), Inches(y))
    c.name = name
    c.line.color.rgb = color
    c.line.width = Pt(0.75)
    c.shadow.inherit = False  # garis tanpa bayangan bawaan tema


def add_badge(slide, name, x, y, number, size=0.44):
    s = add_box(slide, name, (x, y, size, size), fill=BLUE, shape=MSO_SHAPE.OVAL)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, R(str(number), 14, color=WHITE, font=FONT_BOLD))


def add_bar_chart(slide, name, box, categories, values, colors, labels, horizontal=False, font_size=12):
    """Grafik batang asli PowerPoint (datanya bisa diedit). Hanya batang di `labels` yang diberi label."""
    data = CategoryChartData()
    data.categories = categories
    data.add_series("nilai", values)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    x, y, w, h = box
    frame = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
    frame.name = name
    chart = frame.chart
    chart.has_legend = False
    chart.has_title = False
    chart.font.name = FONT
    chart.font.size = Pt(font_size)
    chart.font.color.rgb = INK
    value_axis = chart.value_axis
    value_axis.visible = False  # tanpa sumbu nilai: angka hanya di batang yang disorot
    value_axis.has_major_gridlines = False
    value_axis.minimum_scale = 0
    value_axis.maximum_scale = max(values) * 1.2  # ruang untuk label di ujung batang
    category_axis = chart.category_axis
    category_axis.major_tick_mark = XL_TICK_MARK.NONE
    category_axis.format.line.color.rgb = GRID
    category_axis.tick_labels.font.size = Pt(font_size)
    category_axis.tick_labels.font.color.rgb = INK
    plot = chart.plots[0]
    plot.gap_width = 50
    plot.vary_by_categories = False
    series = plot.series[0]
    for i, color in enumerate(colors):
        series.points[i].format.fill.solid()
        series.points[i].format.fill.fore_color.rgb = color
    for i, (text, color) in labels.items():
        label = series.points[i].data_label
        label.position = XL_LABEL_POSITION.OUTSIDE_END
        add_run(label.text_frame.paragraphs[0], R(text, 16, color=color, font=FONT_BOLD))


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # layout kosong


def add_header(slide, kicker, title):
    add_text(slide, "kicker", (0.6, 0.42, 12.13, 0.3), [P(R(kicker, 11, BLUE, FONT_BOLD))])
    add_text(slide, "judul", (0.6, 0.72, 12.13, 1.0), [P(R(title, 26, font=FONT_BOLD), line=0.95)])


def add_footer(slide, number, text):
    add_line(slide, "garis_footer", 0.6, 6.78, 12.73)
    add_text(slide, "footer_sumber", (0.6, 6.86, 11.2, 0.5), [P(R(text, 9, MUTED))])
    add_text(slide, "footer_halaman", (11.93, 6.86, 0.8, 0.2), [P(R(f"{number} / {N_SLIDES}", 9, MUTED), align=PP_ALIGN.RIGHT)])


def add_panel(slide, name, box, paragraphs):
    """Panel abu-abu muda di sisi kanan slide berisi penjelasan."""
    x, y, w, h = box
    add_box(slide, f"bg_{name}", box, fill=PANEL)
    add_text(slide, name, (x + 0.3, y + 0.3, w - 0.6, h - 0.6), paragraphs)


def heading(text):
    return P(R(text, 11, MUTED, FONT_BOLD), after=6)


def set_notes(slide, lines):
    tf = slide.notes_slide.notes_text_frame
    tf.text = "\n".join(lines)
    for p in tf.paragraphs:
        for run in p.runs:
            run.font.language_id = MSO_LANGUAGE_ID.INDONESIAN


# ---------- slide 1: judul ----------
def slide_title(prs, f):
    s = new_slide(prs)
    kpi = f["kpi"]
    add_box(s, "aksen_kiri", (0, 0, 0.18, 7.5), fill=BLUE)
    add_text(s, "kicker", (0.9, 0.75, 11.5, 0.35), [P(R("EXECUTIVE SUMMARY · UNTUK MANAJEMEN NUSAMART", 12, BLUE, FONT_BOLD))])
    add_text(s, "judul", (0.9, 1.12, 11.5, 0.9), [P(R("Penjualan & customer NusaMart", 40, font=FONT_BOLD))])
    add_text(s, "subjudul", (0.9, 2.05, 11.5, 0.4), [
        P(R("Temuan utama dan rekomendasi · jendela analisis Jan 2017 – Jul 2018 · order delivered", 16, MUTED))
    ])
    add_text(s, "penyusun", (0.9, 2.55, 11.5, 0.35), [P(
        R("Disusun oleh ", 14, MUTED), R(AUTHOR, 14, font=FONT_BOLD), R(f" · {PREPARED_ON}", 14, MUTED),
    )])
    tiles = [
        ("Revenue · tanpa ongkir", f"R$ {idn(kpi['total_revenue'] / 1e6, 2)} juta"),
        ("Order delivered", idn(kpi["total_orders"])),
        ("Customer unik", idn(kpi["total_customers"])),
    ]
    for i, (label, value) in enumerate(tiles, start=1):
        x = 0.9 + (i - 1) * 3.95
        add_box(s, f"bg_kpi_{i}", (x, 3.3, 3.65, 1.6), fill=PANEL)
        add_text(s, f"kpi_{i}", (x + 0.3, 3.45, 3.05, 1.3), [P(R(label, 12, MUTED), after=6), P(R(value, 32, font=FONT_BOLD))],
                 anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, "rekonsiliasi", (0.9, 5.3, 11.5, 0.4), [P(
        R("✓  ", 14, BLUE, "Segoe UI Symbol"),
        R("KPI di notebook (Python), SQL (PostgreSQL), dan dashboard Power BI sudah direkonsiliasi dan identik.", 14),
    )])
    add_footer(s, 1, "NusaMart adalah nama fiktif. Data: Brazilian E-Commerce Public Dataset by Olist (Kaggle), lisensi "
                     "CC BY-NC-SA 4.0; wilayah asli (state Brasil) dipertahankan, mata uang Real Brasil (R$). "
                     "Revenue = Σ harga barang pada order delivered. Sumber angka: results/05_kpi_total_filtered.csv")
    set_notes(s, [
        "Tujuan slide: memperkenalkan cakupan analisis dan membangun kepercayaan pada angka.",
        "NusaMart adalah marketplace fiktif; datanya dataset publik Olist dari Brasil, jadi mata uangnya Real Brasil (R$) "
        "dan wilayahnya state Brasil.",
        "Jendela analisis Jan 2017 – Jul 2018: bulan di kedua ujung data tidak lengkap, dan data Agustus 2018 terpotong "
        "di tengah bulan, jadi dikecualikan.",
        f"KPI jendela analisis: revenue {rp(kpi['total_revenue'])} dari {idn(kpi['total_orders'])} order dan "
        f"{idn(kpi['total_customers'])} customer; AOV {rp(kpi['aov'])}; {idn(kpi['units_sold'])} unit terjual; "
        f"freight burden (proxy) {pct(kpi['freight_burden_pct'])}; repeat rate {pct(kpi['repeat_rate_pct'])}.",
        "Definisi: revenue = jumlah harga barang pada order berstatus delivered, tanpa ongkir; "
        "customer = customer_unique_id.",
        "Rekonsiliasi: angka yang sama muncul di notebook (bagian KPI referensi), sql/05_kpi_reference.sql, dan "
        "halaman Rekonsiliasi di laporan Power BI.",
        "Sumber: results/05_kpi_total_filtered.csv.",
    ])


# ---------- slide 2: masalah bisnis ----------
def slide_problem(prs, f):
    s = new_slide(prs)
    g, one_time, q2c = f["growth"], f["one_time"], f["q2c"]
    add_header(s, "MASALAH BISNIS", "Empat pertanyaan dasar manajemen — dan jawaban singkatnya dari data")
    add_box(s, "bg_konteks", (0.6, 1.85, 12.13, 1.0), line=GREY, dashed=True)
    add_text(s, "konteks", (0.9, 1.98, 11.6, 0.78), [
        P(R("KONTEKS KASUS · SKENARIO, BUKAN TEMUAN DATA", 10, MUTED, FONT_BOLD), after=4),
        P(R("NusaMart (marketplace fiktif) memiliki data transaksi dalam jumlah besar, tetapi manajemen belum dapat "
            "menjawab pertanyaan dasar secara konsisten.", 14, italic=True)),
    ])
    add_text(s, "kolom_pertanyaan", (0.6, 3.07, 4.0, 0.25), [P(R("PERTANYAAN MANAJEMEN", 10, MUTED, FONT_BOLD))])
    add_text(s, "kolom_jawaban", (4.9, 3.07, 7.83, 0.25), [P(R("JAWABAN DARI DATA", 10, BLUE, FONT_BOLD))])
    rows = [
        ("Dari mana pertumbuhan berasal?",
         f"Dari bertambahnya order dan customer: pada Jan–Jul, revenue 2018 naik {pct(g['revenue'], 1)} dibanding 2017, "
         f"sedangkan nilai per order hanya +{pct(g['aov'], 1)}."),
        ("Seberapa loyal customer?",
         f"Sangat rendah: {pct(one_time['pct_customers'], 1)} customer hanya belanja sekali — termasuk customer dengan "
         "belanja terbesar."),
        ("Di mana beban ongkir paling berat?",
         "Di luar São Paulo dan pada kategori barang murah. Ongkir dibayar customer, jadi bebannya dirasakan "
         "langsung oleh pembeli."),
        ("Kategori mana yang bergeser?",
         "health_beauty dan watches_gifts naik ke dua peringkat teratas, sedangkan porsi cool_stuff turun lebih dari "
         "separuh."),
    ]
    for i, (question, answer) in enumerate(rows, start=1):
        y = 3.42 + (i - 1) * 0.8
        add_line(s, f"garis_baris_{i}", 0.6, y, 12.73)
        add_text(s, f"pertanyaan_{i}", (0.6, y + 0.12, 4.0, 0.62), [P(R(question, 15, font=FONT_BOLD))])
        add_text(s, f"jawaban_{i}", (4.9, y + 0.12, 7.83, 0.62), [P(R(answer, 14))])
    add_footer(s, 2, "Sumber: results/02_q1c_fair_period_comparison.csv (periode sebanding Jan–Jul), "
                     "results/03_q4b_order_distribution.csv, results/02_q3_region.csv, "
                     "results/04_q7_high_burden_categories.csv, results/02_q2c_category_fair_period.csv")
    ranks = {c: (int(q2c.loc[c, "rank_2017"]), int(q2c.loc[c, "rank_2018"])) for c in ["health_beauty", "watches_gifts", "cool_stuff"]}
    share = {c: (pct(q2c.loc[c, "share_2017_pct"]), pct(q2c.loc[c, "share_2018_pct"])) for c in ranks}
    set_notes(s, [
        "Kalimat masalah bisnis di kotak atas adalah skenario studi kasus, bukan temuan data. Jawaban di bawahnya "
        "berasal dari data.",
        f"Pertumbuhan, periode sebanding Jan–Jul: revenue +{pct(g['revenue'], 1)}, order +{pct(g['orders'], 1)}, "
        f"customer +{pct(g['customers'], 1)}, AOV +{pct(g['aov'], 1)} — results/02_q1c_fair_period_comparison.csv.",
        f"Loyalitas: {idn(one_time['customers'])} dari {idn(f['kpi']['total_customers'])} customer "
        f"({pct(one_time['pct_customers'], 1)}) hanya order sekali — results/03_q4b_order_distribution.csv.",
        "Beban ongkir: rinciannya di slide 5.",
        "Kategori, periode sebanding Jan–Jul (results/02_q2c_category_fair_period.csv): "
        f"health_beauty naik ke #{ranks['health_beauty'][1]} (porsi {share['health_beauty'][0]} → {share['health_beauty'][1]}); "
        f"watches_gifts dari #{ranks['watches_gifts'][0]} ke #{ranks['watches_gifts'][1]} "
        f"({share['watches_gifts'][0]} → {share['watches_gifts'][1]}); cool_stuff turun dari #{ranks['cool_stuff'][0]} "
        f"ke #{ranks['cool_stuff'][1]} ({share['cool_stuff'][0]} → {share['cool_stuff'][1]}).",
        "Rekomendasi untuk pergeseran kategori ada di catatan slide 6 (rekomendasi lanjutan).",
    ])


# ---------- slide 3: temuan 1, pertumbuhan ----------
def slide_growth(prs, f):
    s = new_slide(prs)
    g, q1c = f["growth"], f["q1c"]
    add_header(s, "TEMUAN 1 · PERTUMBUHAN", "Pertumbuhan datang dari volume — nilai per order hampir tidak naik")
    add_text(s, "judul_grafik", (0.6, 1.9, 7.3, 0.3), [P(R("Perubahan Jan–Jul 2018 dibanding Jan–Jul 2017", 13, font=FONT_BOLD))])
    values = [round(g[c], 1) for c in ["revenue", "orders", "customers", "aov"]]
    add_bar_chart(
        s, "grafik_pertumbuhan", (0.6, 2.3, 7.3, 4.3),
        ["Revenue", "Order", "Customer", "Nilai per order (AOV)"], values,
        colors=[BLUE, BLUE, BLUE, ORANGE],
        labels={0: (f"+{pct(g['revenue'], 1)}", BLUE), 3: (f"+{pct(g['aov'], 1)}", ORANGE)},
        font_size=13,
    )
    add_panel(s, "penjelasan", (8.3, 1.9, 4.43, 4.7), [
        heading("APA ARTINYA"),
        P(R("Revenue, order, dan customer tumbuh hampir sama besar.", 15, font=FONT_BOLD), after=8),
        P(R("Nilai rata-rata per order (AOV) praktis tidak berubah: customer tidak berbelanja lebih banyak per transaksi.", 14), after=8),
        P(R("Pertumbuhan masih bergantung pada tambahan order dan customer baru.", 14), after=18),
        heading("IMPLIKASI"),
        P(R("Perlu pengungkit pertumbuhan kedua: menaikkan nilai per order (lihat rekomendasi).", 14, font=FONT_BOLD)),
    ])
    add_footer(s, 3, "Sumber: results/02_q1c_fair_period_comparison.csv · periode sebanding Jan–Jul, karena data 2018 di "
                     "jendela analisis hanya sampai Juli · AOV = revenue ÷ jumlah order")
    set_notes(s, [
        "Angka lengkap, Jan–Jul 2017 → Jan–Jul 2018 (results/02_q1c_fair_period_comparison.csv): "
        f"revenue {rp(q1c.loc[2017, 'revenue'])} → {rp(q1c.loc[2018, 'revenue'])} (+{pct(g['revenue'], 1)}); "
        f"order {idn(q1c.loc[2017, 'orders'])} → {idn(q1c.loc[2018, 'orders'])} (+{pct(g['orders'], 1)}); "
        f"customer {idn(q1c.loc[2017, 'customers'])} → {idn(q1c.loc[2018, 'customers'])} (+{pct(g['customers'], 1)}); "
        f"AOV {rp(q1c.loc[2017, 'aov'])} → {rp(q1c.loc[2018, 'aov'])} (+{pct(g['aov'], 1)}).",
        "Karena AOV hampir tidak berubah, revenue tumbuh hampir persis sebesar pertumbuhan order dan customer.",
        f"Laju YoY bulanan melambat dari +{pct(f['yoy_jan'], 1)} (Januari 2018) menjadi +{pct(f['yoy_jul'], 1)} "
        "(Juli 2018) — results/02_q1b_yoy.csv.",
        "Hipotesis [High confidence]: angka pertumbuhan yang besar sebagian mencerminkan basis awal 2017 yang masih "
        "kecil, karena YoY terus melambat. Data belum cukup untuk menyimpulkan apakah volume mulai mendatar.",
        "Implikasi: nilai per order adalah pengungkit pertumbuhan berikutnya (rekomendasi prioritas 2).",
    ])


# ---------- slide 4: temuan 2, retensi ----------
def slide_retention(prs, f):
    s = new_slide(prs)
    d1, one_time, cohort, repeat = f["d1"], f["one_time"], f["cohort"], f["repeat"]
    add_header(s, "TEMUAN 2 · RETENSI & KONSENTRASI",
               f"{idn(one_time['pct_customers'])}% customer hanya belanja sekali — bahkan customer bernilai tertinggi jarang kembali")
    add_text(s, "judul_grafik", (0.6, 1.9, 7.3, 0.3), [P(R("Porsi revenue per desil customer (desil 1 = 10% customer teratas)", 13, font=FONT_BOLD))])
    shares = list(f["decile_share"])
    add_bar_chart(
        s, "grafik_desil", (0.6, 2.3, 7.3, 4.3),
        [str(d) for d in f["decile_share"].index], shares,
        colors=[BLUE] + [LIGHT_BLUE] * (len(shares) - 1),
        labels={0: (pct(shares[0]), BLUE)},
        font_size=13,
    )
    add_panel(s, "penjelasan", (8.3, 1.9, 4.43, 4.7), [
        heading("APA ARTINYA"),
        P(R(f"10% customer teratas menyumbang {pct(shares[0])} revenue …", 15, font=FONT_BOLD), after=8),
        P(R(f"… tetapi {pct(d1['revenue_from_one_time_pct'])} revenue kelompok ini berasal dari customer yang hanya "
            "belanja sekali.", 14), after=8),
        P(R("Bukan karena customer belum sempat kembali: customer yang diamati penuh selama setahun pun hampir "
            "semuanya tidak kembali.", 14), after=18),
        heading("IMPLIKASI"),
        P(R("Jadikan pembelian pertama — terutama yang bernilai besar — pemicu program pembelian kedua.", 14, font=FONT_BOLD)),
    ])
    add_footer(s, 4, "Sumber: results/03_q4b_order_distribution.csv, results/03_q5a_decile.csv, "
                     "results/03_q5c_decile_repeat.csv, results/03_q4d_cohort_12m_repeat.csv · customer = customer_unique_id")
    set_notes(s, [
        f"Repeat rate {pct(repeat['repeat_rate_pct'])}: {idn(repeat['repeat_customers'])} dari "
        f"{idn(repeat['total_customers'])} customer order dua kali atau lebih; {idn(one_time['customers'])} customer "
        f"({pct(one_time['pct_customers'], 1)}) hanya order sekali — results/03_q4a_repeat_rate.csv, "
        "results/03_q4b_order_distribution.csv.",
        f"Bukan karena customer belum sempat kembali: dari {idn(cohort['customers'])} customer yang order pertamanya "
        f"Jan–Jun 2017 (diamati penuh 12 bulan), hanya {idn(cohort['repeat_12m'])} ({pct(cohort['repeat_12m_pct'])}) "
        "yang order lagi di hari lain dalam 12 bulan — results/03_q4d_cohort_12m_repeat.csv.",
        f"Customer returning hanya menyumbang {pct(f['returning_share'])} revenue di jendela analisis "
        f"({pct(f['returning_share_2018'])} di Jan–Jul 2018) — results/03_q4c_new_vs_returning.csv.",
        f"Konsentrasi: desil 1 ({idn(d1['customers'])} customer) = {pct(shares[0])} revenue; 20% customer teratas = "
        f"{pct(f['top20'])}; separuh customer terbawah = {pct(f['bottom_half'])} — results/03_q5a_decile.csv.",
        f"Desil 1: rata-rata belanja {rp(d1['avg_revenue_per_customer'])} per customer, repeat rate "
        f"{pct(d1['repeat_rate_pct'])}, dan {pct(d1['revenue_from_one_time_pct'])} revenue-nya dari pembeli sekali — "
        "results/03_q5c_decile_repeat.csv.",
        "Hipotesis [Low confidence]: banyak kategori besar jarang dibeli ulang, dan belum ada program platform yang "
        "mengajak customer kembali. Perlu data repeat rate per kategori dan data CRM untuk membuktikannya.",
    ])


# ---------- slide 5: temuan 3, beban ongkir ----------
def slide_freight(prs, f):
    s = new_slide(prs)
    kpi, q3, big = f["kpi"], f["q3"], f["big_states"]
    top_state = big.index[-1]  # freight burden tertinggi di antara state bervolume besar
    add_header(s, "TEMUAN 3 · BEBAN ONGKIR (PROXY)",
               "Di luar São Paulo dan pada barang murah, customer menanggung ongkir lebih berat")
    add_text(s, "judul_grafik", (0.6, 1.9, 7.3, 0.3), [P(R("Freight burden (proxy) per state · ongkir ÷ harga barang", 13, font=FONT_BOLD))])
    # urutkan naik: batang pertama digambar paling bawah, jadi beban tertinggi berada di atas
    rows = [(st, v) for st, v in big["freight_burden_pct"].items()] + [("Rata-rata nasional", kpi["freight_burden_pct"])]
    rows.sort(key=lambda r: r[1])
    names = ["SP (São Paulo)" if st == "SP" else st for st, _ in rows]
    colors = [BLUE if st == "SP" else ORANGE if st == top_state else GREY if st.startswith("Rata") else LIGHT_ORANGE for st, _ in rows]
    index = {st: i for i, (st, _) in enumerate(rows)}
    add_bar_chart(
        s, "grafik_state", (0.6, 2.25, 7.3, 4.4), names, [v for _, v in rows], colors,
        labels={index["SP"]: (pct(q3.loc["SP", "freight_burden_pct"]), BLUE),
                index[top_state]: (pct(q3.loc[top_state, "freight_burden_pct"]), ORANGE)},
        horizontal=True, font_size=11,
    )
    # state besar lain selalu di atas SP, tetapi sebagian hampir sama dengan rata-rata nasional: jangan disebut "jauh"
    others = big.drop("SP")
    near = [st for st, v in others["freight_burden_pct"].items() if abs(v - kpi["freight_burden_pct"]) <= NEAR_AVG_PP]
    near_text = f"dari hampir sama dengan rata-rata nasional ({' dan '.join(near)}) " if near else ""
    add_panel(s, "penjelasan", (8.3, 1.9, 4.43, 4.75), [
        heading("PER STATE"),
        P(R(f"São Paulo, pasar terbesar, punya beban ongkir terendah. State besar lainnya lebih tinggi: {near_text}"
            f"hingga {pct(q3.loc[top_state, 'freight_burden_pct'])} di {top_state}.", 13), after=14),
        heading("PER KATEGORI"),
        P(R(f"Kategori besar dengan freight burden di atas rata-rata menyumbang {pct(f['q7_share'])} revenue.", 14,
            font=FONT_BOLD), after=6),
        P(R("Hampir semuanya barang murah, kecuali office_furniture yang ongkirnya mahal.", 13), after=14),
        heading("IMPLIKASI · HIPOTESIS (KEYAKINAN RENDAH)"),
        P(R("Ongkir yang berat berpotensi menahan pembelian di luar SP dan pembelian barang murah — perlu diuji.", 14,
            font=FONT_BOLD)),
    ])
    add_footer(s, 5, "Freight burden (proxy) = Σ ongkir ÷ Σ harga barang; ongkir dibayar customer — bukan profit dan bukan "
                     f"biaya perusahaan. Rata-rata nasional {pct(kpi['freight_burden_pct'])}. Hanya state dengan "
                     f"≥ {idn(MIN_STATE_ORDERS)} order yang ditampilkan ({len(big)} dari {len(q3)} state, "
                     f"{pct(f['big_order_share'])} order) agar angkanya stabil. Sumber: results/02_q3_region.csv, "
                     "results/04_q7_high_burden_categories.csv, results/04_q6_volume_vs_value.csv")
    by_burden = big.sort_values("freight_burden_pct", ascending=False)
    sp = q3.loc["SP"]
    set_notes(s, [
        "Freight burden (proxy) = Σ ongkir ÷ Σ harga barang. Ongkir di dataset ini dibayar customer, jadi angka ini "
        "mengukur beban ongkir customer relatif terhadap harga barang — bukan profit dan bukan biaya perusahaan. "
        f"Rata-rata nasional {pct(kpi['freight_burden_pct'])}.",
        f"State dengan ≥ {idn(MIN_STATE_ORDERS)} order ({len(big)} dari {len(q3)} state, {pct(f['big_order_share'])} "
        "order), dari beban tertinggi: "
        + ", ".join(f"{st} {pct(v)}" for st, v in by_burden["freight_burden_pct"].items())
        + " — results/02_q3_region.csv.",
        f"SP: {pct(sp['revenue_share_pct'])} revenue dan {pct(f['sp_order_share'])} order, dengan freight burden "
        f"{pct(sp['freight_burden_pct'])} — terendah dari {len(q3)} state.",
        "State besar lainnya semuanya di atas SP, tetapi tidak semuanya jauh di atas rata-rata: "
        + " dan ".join(f"{st} {pct(others.loc[st, 'freight_burden_pct'])}" for st in near)
        + f" hampir sama dengan rata-rata nasional {pct(kpi['freight_burden_pct'])}, sedangkan yang tertinggi "
        f"{top_state} {pct(q3.loc[top_state, 'freight_burden_pct'])}.",
        "State kecil tidak dijadikan contoh utama karena ordernya sedikit sehingga angkanya mudah berfluktuasi. "
        "Beban tertinggi di antaranya: "
        + ", ".join(f"{st} {pct(r.freight_burden_pct)} ({idn(r.orders)} order)" for st, r in f["small_states"].head(6).iterrows())
        + ".",
        f"Kategori (results/04_q7_high_burden_categories.csv, results/04_q6_volume_vs_value.csv): {len(f['q7'])} kategori "
        f"dengan revenue di atas rata-rata kategori dan freight burden di atas {pct(kpi['freight_burden_pct'])}; total "
        f"{rp(f['q7_revenue'])} ({pct(f['q7_share'])} revenue).",
        f"Semua {len(f['cheap'])} kategori lainnya (selain office_furniture) berharga di bawah rata-rata "
        f"{rp(f['avg_price'])} per unit: ongkirnya biasa, harganya murah. office_furniture: ongkir ≈ "
        f"{rp(f['freight_unit_office'])} per unit, sekitar {f['freight_unit_office'] / f['freight_unit_avg']:.0f}× "
        f"rata-rata ≈ {rp(f['freight_unit_avg'])} (freight burden × harga rata-rata per unit).",
        "Implikasi, hipotesis [Low confidence]: ongkir yang berat berpotensi menahan pembelian di luar SP dan pembelian "
        "barang murah. Dataset hanya berisi order yang sudah terjadi (tanpa data kunjungan atau keranjang yang "
        "ditinggalkan), jadi dampaknya belum bisa diukur; uji lewat gratis ongkir multi-item (prioritas 3) dan pilot "
        "subsidi ongkir per state (rekomendasi lanjutan 4).",
        "Hipotesis [Medium confidence]: seller terkonsentrasi di sekitar SP sehingga jarak kirim ke SP pendek; lokasi "
        "seller berada di luar scope data.",
    ])


# ---------- slide 6: rekomendasi ----------
def slide_recommendations(prs, f):
    s = new_slide(prs)
    kpi, cohort, q3, q2c = f["kpi"], f["cohort"], f["q3"], f["q2c"]
    add_header(s, "REKOMENDASI", "Tiga prioritas: ajak customer kembali, naikkan nilai per order, ringankan beban ongkir")
    add_text(s, "kriteria", (0.6, 1.9, 12.13, 0.35), [P(
        R("Kriteria prioritas:  ", 14, BLUE, FONT_BOLD),
        R("ukuran masalah  ×  kendali platform  ×  kemudahan diuji lewat eksperimen", 14),
    )])
    cards = [
        ("Program pembelian kedua",
         "Voucher platform, notifikasi dan email, serta rekomendasi produk berdasarkan pembelian pertama; diuji "
         "dengan kelompok kontrol.",
         "CRM / Marketing Manager",
         f"Customer baru yang kembali dalam 12 bulan: {pct(cohort['repeat_12m_pct'])}",
         "PAGAR", "AOV kelompok uji tidak turun dibanding kontrol"),
        ("Naikkan nilai per order",
         "Rekomendasi produk pelengkap, voucher untuk keranjang multi-item, dan ambang gratis ongkir yang disubsidi "
         "platform.",
         "Head of Commercial",
         f"Nilai per order (AOV): {rp(kpi['aov'])}",
         "PAGAR", "jumlah order tidak turun"),
        ("Ringankan beban ongkir",
         "Rekomendasi “produk lain dari seller yang sama” dan gratis ongkir multi-item bersubsidi; subsidi terarah "
         "untuk office_furniture.",
         "Category Manager, Operations / Logistics",
         f"Freight burden (proxy): {pct(kpi['freight_burden_pct'])}",
         "METRIK HASIL", "pertumbuhan order di kategori/state sasaran"),
    ]
    width, gap, top = 3.843, 0.3, 2.42
    for i, (title, levers, owner, metric, guard_label, guard) in enumerate(cards, start=1):
        x = 0.6 + (i - 1) * (width + gap)
        add_box(s, f"bg_kartu_{i}", (x, top, width, 4.24), fill=PANEL)
        add_badge(s, f"badge_{i}", x + 0.3, top + 0.26, i)
        add_text(s, f"kartu_{i}_judul", (x + 0.85, top + 0.26, width - 1.05, 0.44), [P(R(title, 16, font=FONT_BOLD))],
                 anchor=MSO_ANCHOR.MIDDLE)
        # posisi bagian dibuat tetap supaya baris "Pemilik", "Metrik", dan "Pagar" sejajar di ketiga kartu
        for part, (label, text, y, h, font) in enumerate([
            ("TUAS PLATFORM", levers, 3.32, 1.3, FONT),
            ("PEMILIK", owner, 4.68, 0.66, FONT),
            ("METRIK (BASELINE)", metric, 5.4, 0.66, FONT_BOLD),
        ], start=1):
            add_text(s, f"kartu_{i}_bagian_{part}", (x + 0.3, y, width - 0.6, h), [
                P(R(label, 10, MUTED, FONT_BOLD), after=3),
                P(R(text, 13, font=font)),
            ])
        add_text(s, f"kartu_{i}_pagar", (x + 0.3, 6.12, width - 0.6, 0.44), [P(R(f"{guard_label}  ", 10, MUTED, FONT_BOLD), R(guard, 12))])
    add_footer(s, 6, "Semua tuas dikendalikan platform, karena di marketplace harga dan ongkir ditentukan seller dan kurir. "
                     "Dampak ditulis sebagai arah: biaya voucher dan subsidi belum bisa dihitung karena dataset tidak memuat "
                     "cost. Baseline: results/03_q4d_cohort_12m_repeat.csv, results/05_kpi_total_filtered.csv · "
                     "detail: docs/insights_and_recommendations.md")
    top_cheap = f["cheap"].head(3)
    pilots = " atau ".join(f"{st} (freight burden {pct(q3.loc[st, 'freight_burden_pct'])}, {idn(q3.loc[st, 'orders'])} order)"
                           for st in PILOT_STATES)
    set_notes(s, [
        "Kriteria prioritas: ukuran masalah × kendali platform × kemudahan diuji. Urutan ini adalah penilaian, bukan "
        "hasil perhitungan.",
        "1) Program pembelian kedua — CRM / Marketing Manager, dibantu Category Manager untuk pemetaan produk pelengkap. "
        "Metrik: customer baru yang order lagi di hari lain dalam 12 bulan (baseline kohort Jan–Jun 2017: "
        f"{pct(cohort['repeat_12m_pct'])}), porsi revenue returning (baseline Jan–Jul 2018: "
        f"{pct(f['returning_share_2018'])}), dan repeat rate desil 1 (baseline {pct(f['d1']['repeat_rate_pct'])}). "
        "Pagar: AOV kelompok uji tidak turun dibanding kontrol, supaya voucher tidak sekadar menjadi diskon.",
        f"2) Nilai per order — Head of Commercial. Metrik: AOV (baseline {rp(kpi['aov'])}) dan unit per order (baseline "
        f"{idn(f['units_per_order'], 2)}), dibanding kelompok kontrol. Pagar: jumlah order kelompok uji tidak turun, "
        "karena ambang gratis ongkir bisa membuat customer menunda pembelian kecil.",
        "3) Beban ongkir — Category Manager bersama Operations / Logistics. Mulai dari kategori murah dengan freight "
        "burden tertinggi: "
        + ", ".join(f"{c} {pct(v)}" for c, v in zip(top_cheap["category_en"], top_cheap["freight_burden_pct"]))
        + "; untuk office_furniture, uji subsidi ongkir terarah. Metrik: freight burden per kategori (baseline "
        f"rata-rata {pct(kpi['freight_burden_pct'])}) dan unit per order. Metrik hasil: pertumbuhan order di kategori "
        "dan state sasaran dibanding kontrol; baseline order per kategori ada di results/02_q2_category.csv dan per "
        "state di results/02_q3_region.csv.",
        f"Rekomendasi lanjutan 4: pilot subsidi ongkir platform di state berongkir tinggi yang volumenya cukup, misalnya "
        f"{pilots}, dengan state lain sebagai pembanding. Metrik: order dan revenue state pilot, serta porsi revenue "
        f"di luar SP (baseline {pct(f['outside_sp'])}).",
        f"Rekomendasi lanjutan 5: investigasi penurunan cool_stuff (porsi {pct(q2c.loc['cool_stuff', 'share_2017_pct'])} → "
        f"{pct(q2c.loc['cool_stuff', 'share_2018_pct'])}) sebelum bertindak; tambah pasokan di kategori yang tumbuh "
        "(health_beauty, watches_gifts, baby); wajibkan periode sebanding di laporan tahunan.",
        "Risiko: voucher, subsidi ongkir, dan potongan komisi adalah biaya platform yang belum bisa diukur karena "
        "dataset tidak memuat cost; karena itu dampak ditulis sebagai arah.",
    ])


# ---------- slide 7: langkah berikutnya ----------
def slide_next_steps(prs, f):
    s = new_slide(prs)
    add_header(s, "LANGKAH BERIKUTNYA", "Uji dulu, ukur terhadap baseline, lalu lengkapi data yang belum ada")
    add_text(s, "judul_langkah", (0.6, 1.9, 7.3, 0.3), [P(R("URUTAN LANGKAH", 11, MUTED, FONT_BOLD))])
    steps = [
        "Rancang eksperimen untuk prioritas 1 dan 2: kelompok uji vs kelompok kontrol, dengan metrik dan baseline "
        "di slide sebelumnya.",
        "Kerjakan analisis lanjutan dari data yang sudah ada: repeat rate per kategori pembelian pertama, komposisi "
        "kategori per desil, serta ongkir per item pada order satu item vs banyak item.",
        "Lengkapi data yang belum ada: cost dan komisi (untuk menghitung biaya program), histori promo, lokasi seller, "
        "serta berat dan dimensi produk.",
        "Pantau KPI secara rutin lewat dashboard Power BI.",
    ]
    for i, step in enumerate(steps, start=1):
        y = 2.35 + (i - 1) * 0.92
        add_badge(s, f"badge_{i}", 0.6, y, i, size=0.4)
        add_text(s, f"langkah_{i}", (1.2, y + 0.02, 6.7, 0.82), [P(R(step, 14))])
    add_panel(s, "keterbatasan", (8.3, 1.9, 4.43, 3.95), [
        heading("KETERBATASAN"),
        P(R("Tidak ada data cost, profit, atau komisi: freight burden hanya proxy, dan dampak finansial rekomendasi "
            "belum bisa dihitung.", 13), after=8, bullet=True),
        P(R("Jendela analisis Jan 2017 – Jul 2018; data Agustus 2018 terpotong sehingga dikecualikan.", 13), after=8,
          bullet=True),
        P(R("Penyebab masih berupa hipotesis dan korelasi bukan sebab-akibat, jadi rekomendasi perlu diuji dulu.", 13),
          bullet=True),
    ])
    add_line(s, "garis_link", 0.6, 5.98, 12.73)
    # Saat PowerPoint mengekspor PDF, hyperlink yang menjadi teks terakhir di text box tidak ikut menjadi link.
    # Karena itu setiap link diberi spasi biasa di belakangnya.
    add_text(s, "link_dashboard", (0.6, 6.08, 12.13, 0.28), [P(
        R("Dashboard interaktif (Publish to web): ", 12, font=FONT_BOLD),
        R("buka dashboard Power BI", 12, BLUE, link=DASHBOARD_URL),
        R(" ", 12),
    )])
    add_text(s, "link_repo", (0.6, 6.4, 12.13, 0.28), [P(
        R("Kode & analisis (GitHub): ", 12, font=FONT_BOLD),
        R(REPO_URL.removeprefix("https://"), 12, BLUE, link=REPO_URL),
        R(" ", 12),
    )])
    add_footer(s, 7, "Detail temuan, hipotesis, dan rekomendasi: docs/insights_and_recommendations.md · semua angka di "
                     "deck ini berasal dari results/ (hasil query SQL)")
    set_notes(s, [
        "Langkahnya berurutan dan sengaja tanpa tanggal: eksperimen dulu, karena semua penyebab masih hipotesis dan "
        "biaya program belum bisa diukur tanpa data cost.",
        "Analisis lanjutan yang bisa langsung dikerjakan dengan data yang sudah ada: repeat rate per kategori pembelian "
        "pertama (Insight 2), komposisi kategori per desil (Insight 3), dan ongkir per item pada order satu item vs "
        "banyak item (Insight 4).",
        "Data yang dibutuhkan: cost dan komisi (untuk menghitung biaya voucher dan subsidi), histori promo, lokasi "
        "seller, serta berat dan dimensi produk (ada di file mentah, belum dimodelkan).",
        "Keterbatasan lain: dokumentasi dataset tidak menjelaskan apakah isinya seluruh transaksi atau sampel, jadi "
        "volume absolut sebaiknya tidak dibaca sebagai ukuran pasar.",
        "Semua query, angka pendukung, dan dashboard bisa dibuka lewat link di slide ini.",
    ])


def main():
    facts = load_facts()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9
    for build in [slide_title, slide_problem, slide_growth, slide_retention, slide_freight,
                  slide_recommendations, slide_next_steps]:
        build(prs, facts)
    props = prs.core_properties
    props.title = "NusaMart — Executive Summary"
    props.subject = "Temuan utama dan rekomendasi, jendela analisis Jan 2017 – Jul 2018"
    props.author = props.last_modified_by = AUTHOR
    props.keywords = "NusaMart; Olist; e-commerce; executive summary"
    props.comments = "Dibuat oleh src/build_executive_summary.py dari file di results/"
    props.created = props.modified = datetime.now()
    props.revision = 1
    OUT.parent.mkdir(exist_ok=True)
    prs.save(OUT)
    print(f"Tersimpan: {OUT.relative_to(ROOT)} ({len(prs.slides)} slide)")


if __name__ == "__main__":
    main()
